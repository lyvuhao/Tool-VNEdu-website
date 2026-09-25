"""Nhập điểm từ file Excel/CSV: đọc bảng, đoán cột, khớp học sinh và lập kế hoạch điểm chờ.

Module không phụ thuộc giao diện. Kết quả là danh sách `ImportPlanRow`; giao diện
(`ui/excel_import.py`) cho giáo viên xem lại rồi đưa các dòng đã chọn vào hàng "Chờ ghi".
Việc ghi lên VNEDU vẫn dùng nút "GHI ĐIỂM LÊN WEB" như với điểm đọc bằng giọng nói.
"""

from __future__ import annotations

import csv
import io
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from .matching import _token_sort_ratio
from .models import ScoreStudentRow
from .scores import _format_score_value, parse_manual_score_text
from .voice.phonetics import _normalize_diacritic_text

SUPPORTED_EXTENSIONS = (".xlsx", ".xlsm", ".csv", ".txt")
HEADER_SCAN_ROWS = 20
FUZZY_MIN_SCORE = 88
FUZZY_MIN_MARGIN = 6

# Trạng thái một dòng trong file sau khi khớp
MATCH_CODE = "code"
MATCH_EXACT = "exact"
MATCH_ASCII = "ascii"
MATCH_ORDER = "order"
MATCH_FUZZY = "fuzzy"
NO_MATCH = "no_match"
EMPTY_SCORE = "empty_score"
INVALID_SCORE = "invalid_score"
DUPLICATE_SOURCE = "duplicate_source"
SAME_AS_CURRENT = "same_as_current"

# Dòng khớp chắc chắn -> được chọn sẵn; còn lại giáo viên tự tích nếu đồng ý.
CONFIDENT_MATCHES = {MATCH_CODE, MATCH_EXACT, MATCH_ASCII}

_NAME_PREFIXES = ("ho va ten", "ho ten", "ten hoc sinh", "hoc sinh")
_SURNAME_HEADERS = ("ho", "ho dem", "ho va ten dem", "ho va dem", "ho lot", "ho va chu lot")
_GIVEN_HEADERS = ("ten",)
_CODE_PREFIXES = ("ma hoc sinh", "ma hs", "ma dinh danh", "ma so")
_CODE_HEADERS = ("ma", "id")
_SCORE_HINTS = ("diem", "tx", "gk", "ck", "giua ky", "cuoi ky", "dtb", "tbm", "kiem tra", "ktra", "ddg")


def _is_surname_header(key: str) -> bool:
    return key in _SURNAME_HEADERS


def _is_name_header(key: str) -> bool:
    return not _is_surname_header(key) and key.startswith(_NAME_PREFIXES)


def _is_code_header(key: str) -> bool:
    return key in _CODE_HEADERS or key.startswith(_CODE_PREFIXES)


def _is_score_header(key: str) -> bool:
    return bool(key) and any(hint in key.split() or key.startswith(hint) for hint in _SCORE_HINTS)


class ImportFileError(Exception):
    """Lỗi đọc file nhập điểm (thông điệp hiển thị cho người dùng)."""


@dataclass
class ImportedTable:
    """Bảng đọc từ file: tiêu đề + các dòng dữ liệu (số dòng tính theo file, bắt đầu từ 1)."""

    source_path: Path
    sheet_name: str
    sheet_names: list[str]
    header_row: int
    headers: list[str]
    rows: list[tuple[int, list[str]]] = field(default_factory=list)


@dataclass
class ColumnChoice:
    """Cột dùng để khớp. `given` chỉ dùng khi họ và tên nằm ở hai cột riêng."""

    name: int | None = None
    given: int | None = None
    code: int | None = None
    score: int | None = None


@dataclass
class ImportPlanRow:
    source_row: int
    source_name: str
    source_code: str
    raw_score: str
    score: str
    status: str
    detail: str
    row_key: str = ""
    student_name: str = ""
    current_score: str = ""
    pending_score: str = ""
    selected: bool = False

    @property
    def can_apply(self) -> bool:
        return bool(self.row_key and self.score) and self.status not in (DUPLICATE_SOURCE, SAME_AS_CURRENT)


# ---------------------------------------------------------------------------
# Đọc file
# ---------------------------------------------------------------------------

def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".10g")
    return " ".join(str(value).split())


def _trim_rows(raw_rows: Iterable[Sequence[object]]) -> list[list[str]]:
    rows = [[_cell_text(cell) for cell in row] for row in raw_rows]
    width = max((len(row) for row in rows), default=0)
    # Bỏ các cột trống ở cuối
    while width and all(len(row) < width or not row[width - 1] for row in rows):
        width -= 1
    return [(row + [""] * width)[:width] for row in rows]


def _read_csv_rows(path: Path) -> list[list[str]]:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1258", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("latin-1")
    # CSV "Windows tiếng Việt" (cp1258) lưu dấu thanh dạng tổ hợp -> chuẩn hoá về dạng dựng sẵn.
    text = unicodedata.normalize("NFC", text)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _trim_rows(csv.reader(io.StringIO(text), dialect))


def list_sheet_names(path: Path) -> list[str]:
    """Tên các sheet (file CSV coi như một sheet)."""

    path = Path(path)
    if path.suffix.lower() in (".csv", ".txt"):
        return [path.stem]
    workbook = _open_workbook(path)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def _open_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportFileError(
            "Chưa cài thư viện openpyxl để đọc file Excel.\n"
            "Cài bằng lệnh: pip install openpyxl — hoặc lưu file dưới dạng CSV rồi nhập lại."
        ) from error
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - file hỏng/khác định dạng
        raise ImportFileError(f"Không đọc được file Excel: {error}") from error


def read_score_table(path: Path, sheet_name: str | None = None) -> ImportedTable:
    """Đọc một sheet (hoặc file CSV) và tự tìm dòng tiêu đề."""

    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise ImportFileError("File .xls (Excel 97-2003) chưa được hỗ trợ. Hãy mở bằng Excel và lưu lại dạng .xlsx.")
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ImportFileError(f"Chỉ hỗ trợ file {', '.join(SUPPORTED_EXTENSIONS)}.")
    if not path.is_file():
        raise ImportFileError(f"Không tìm thấy file: {path}")

    if suffix in (".csv", ".txt"):
        sheet_names = [path.stem]
        chosen = path.stem
        rows = _read_csv_rows(path)
    else:
        workbook = _open_workbook(path)
        try:
            sheet_names = list(workbook.sheetnames)
            chosen = sheet_name if sheet_name in sheet_names else sheet_names[0]
            rows = _trim_rows(workbook[chosen].iter_rows(values_only=True))
        finally:
            workbook.close()

    numbered = [(index, row) for index, row in enumerate(rows, start=1) if any(row)]
    if not numbered:
        raise ImportFileError("File không có dữ liệu.")
    header_row, headers = _find_header(numbered)
    data_rows = [(number, row) for number, row in numbered if number > header_row]
    if header_row == 0:
        headers = [f"Cột {_column_letter(i)}" for i in range(len(rows[0]) if rows else 0)]
    return ImportedTable(
        source_path=path,
        sheet_name=chosen,
        sheet_names=sheet_names,
        header_row=header_row,
        headers=headers,
        rows=data_rows,
    )


def _column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _header_key(text: str) -> str:
    return _normalize_diacritic_text(text)


def _header_score(row: Sequence[str]) -> int:
    keys = [_header_key(cell) for cell in row]
    score = 0
    if any(_is_name_header(key) or _is_surname_header(key) for key in keys):
        score += 3
    if any(key in _GIVEN_HEADERS for key in keys):
        score += 1
    if any(_is_code_header(key) for key in keys):
        score += 1
    if any(_is_score_header(key) for key in keys):
        score += 2
    if any(key in ("stt", "tt") for key in keys):
        score += 1
    return score


def _find_header(numbered: list[tuple[int, list[str]]]) -> tuple[int, list[str]]:
    best_number, best_row, best_score = 0, [], 0
    for number, row in numbered[:HEADER_SCAN_ROWS]:
        score = _header_score(row)
        if score > best_score:
            best_number, best_row, best_score = number, row, score
    if best_score >= 3:
        return best_number, [cell or f"Cột {_column_letter(i)}" for i, cell in enumerate(best_row)]
    return 0, []


# ---------------------------------------------------------------------------
# Đoán cột
# ---------------------------------------------------------------------------

def guess_columns(table: ImportedTable) -> ColumnChoice:
    """Đoán cột họ tên / tên / mã HS / điểm từ tiêu đề, không có tiêu đề thì đoán theo nội dung."""

    choice = ColumnChoice()
    keys = [_header_key(header) for header in table.headers]
    if table.header_row:
        for index, key in enumerate(keys):
            if choice.name is None and _is_name_header(key):
                choice.name = index
            elif choice.code is None and _is_code_header(key):
                choice.code = index
        if choice.name is None:
            surname = next((i for i, key in enumerate(keys) if _is_surname_header(key)), None)
            given = next((i for i, key in enumerate(keys) if key in _GIVEN_HEADERS), None)
            if surname is not None:
                choice.name = surname
                choice.given = given
            elif given is not None:
                choice.name = given
        for index, key in enumerate(keys):
            if index in (choice.name, choice.given, choice.code):
                continue
            if _is_score_header(key):
                choice.score = index
                break

    columns = range(len(table.headers))
    if choice.name is None:
        choice.name = max(columns, key=lambda i: _text_column_ratio(table, i), default=None)
    if choice.score is None:
        candidates = [i for i in columns if i not in (choice.name, choice.given, choice.code)]
        best = max(candidates, key=lambda i: _score_column_ratio(table, i), default=None)
        if best is not None and _score_column_ratio(table, best) > 0.5:
            choice.score = best
    return choice


def _column_values(table: ImportedTable, index: int) -> list[str]:
    return [row[index] for _number, row in table.rows if index < len(row) and row[index]]


def _text_column_ratio(table: ImportedTable, index: int) -> float:
    values = _column_values(table, index)
    if not values:
        return 0.0
    names = sum(1 for value in values if len(value.split()) >= 2 and any(ch.isalpha() for ch in value))
    return names / len(values)


def _score_column_ratio(table: ImportedTable, index: int) -> float:
    values = _column_values(table, index)
    if not values or _looks_like_ordinal(values):
        return 0.0
    return sum(1 for value in values if parse_manual_score_text(value) is not None) / len(values)


def _looks_like_ordinal(values: Sequence[str]) -> bool:
    """Cột STT (1, 2, 3, ...) cũng là số 0-10 nhưng không phải điểm."""
    if len(values) < 2 or not all(value.isdigit() for value in values):
        return False
    numbers = [int(value) for value in values]
    return numbers == list(range(numbers[0], numbers[0] + len(numbers)))


# ---------------------------------------------------------------------------
# Khớp học sinh
# ---------------------------------------------------------------------------

def _exact_key(text: str) -> str:
    return " ".join(unicodedata.normalize("NFC", str(text)).casefold().split())


def _cell(row: Sequence[str], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return row[index].strip()


def build_import_plan(
    table: ImportedTable,
    columns: ColumnChoice,
    students: Sequence[ScoreStudentRow],
) -> list[ImportPlanRow]:
    """Khớp từng dòng trong file với học sinh của lớp đang mở trên VNEDU."""

    if columns.name is None and columns.code is None:
        raise ImportFileError("Hãy chọn cột họ tên hoặc cột mã học sinh.")
    if columns.score is None:
        raise ImportFileError("Hãy chọn cột điểm.")

    ordered = sorted(students, key=lambda row: row.row_index)
    by_code: dict[str, list[ScoreStudentRow]] = {}
    by_exact: dict[str, list[ScoreStudentRow]] = {}
    by_ascii: dict[str, list[ScoreStudentRow]] = {}
    for student in ordered:
        if student.student_code.strip():
            by_code.setdefault(student.student_code.strip().casefold(), []).append(student)
        by_exact.setdefault(_exact_key(student.student_name), []).append(student)
        by_ascii.setdefault(_normalize_diacritic_text(student.student_name), []).append(student)

    plan: list[ImportPlanRow] = []
    used_keys: dict[str, int] = {}
    taken: set[str] = set()
    for source_row, row in table.rows:
        name = _cell(row, columns.name)
        given = _cell(row, columns.given)
        if given:
            name = f"{name} {given}".strip()
        code = _cell(row, columns.code)
        raw_score = _cell(row, columns.score)
        if not name and not code:
            continue
        item = ImportPlanRow(source_row=source_row, source_name=name, source_code=code, raw_score=raw_score,
                             score="", status=NO_MATCH, detail="")

        student, status, detail = _match_student(name, code, ordered, by_code, by_exact, by_ascii, taken)
        item.status, item.detail = status, detail
        if student is not None:
            taken.add(student.row_key)
            item.row_key = student.row_key
            item.student_name = student.student_name
            item.current_score = student.current_score
            item.pending_score = student.pending_score

        score_value = parse_manual_score_text(raw_score)
        if not raw_score:
            item.status, item.detail = (EMPTY_SCORE, "Ô điểm trống — bỏ qua.") if student is not None else (item.status, item.detail)
        elif score_value is None:
            item.status, item.detail = INVALID_SCORE, f"Điểm `{raw_score}` không hợp lệ (0-10)."
        else:
            item.score = _format_score_value(score_value)

        if item.row_key and item.score and item.status not in (EMPTY_SCORE, INVALID_SCORE):
            if item.row_key in used_keys:
                item.status = DUPLICATE_SOURCE
                item.detail = f"Học sinh này đã có ở dòng {used_keys[item.row_key]} của file."
            else:
                used_keys[item.row_key] = source_row
                if not item.pending_score and _same_score(item.score, item.current_score):
                    item.status = SAME_AS_CURRENT
                    item.detail = "Trùng điểm đang có trên VNEDU — không cần ghi."
                elif item.pending_score and not _same_score(item.score, item.pending_score):
                    item.detail += f" Thay điểm chờ {item.pending_score}."
                elif item.current_score and not _same_score(item.score, item.current_score):
                    item.detail += f" Sẽ ghi đè điểm hiện tại {item.current_score}."
        item.selected = item.can_apply and item.status in CONFIDENT_MATCHES
        plan.append(item)
    return plan


def _same_score(left: str, right: str) -> bool:
    try:
        return abs(float(str(left).replace(",", ".")) - float(str(right).replace(",", "."))) < 1e-9
    except ValueError:
        return False


def _match_student(
    name: str,
    code: str,
    ordered: Sequence[ScoreStudentRow],
    by_code: dict[str, list[ScoreStudentRow]],
    by_exact: dict[str, list[ScoreStudentRow]],
    by_ascii: dict[str, list[ScoreStudentRow]],
    taken: set[str],
) -> tuple[ScoreStudentRow | None, str, str]:
    if code:
        hits = by_code.get(code.casefold(), [])
        if len(hits) == 1:
            return hits[0], MATCH_CODE, "Khớp mã học sinh."
    if not name:
        return None, NO_MATCH, "Không tìm thấy mã học sinh trong lớp."

    for index, status, label in ((by_exact, MATCH_EXACT, "Khớp họ tên."), (by_ascii, MATCH_ASCII, "Khớp họ tên (bỏ dấu).")):
        key = _exact_key(name) if status == MATCH_EXACT else _normalize_diacritic_text(name)
        hits = index.get(key, [])
        if len(hits) == 1:
            return hits[0], status, label
        if len(hits) > 1:
            # Trùng tên trong lớp: lấy bạn đầu tiên (theo thứ tự trên VNEDU) chưa được dòng nào trước đó khớp.
            free = [hit for hit in hits if hit.row_key not in taken]
            if free:
                return free[0], MATCH_ORDER, f"Lớp có {len(hits)} bạn trùng tên — khớp theo thứ tự, cần kiểm tra."
            return None, NO_MATCH, f"Lớp có {len(hits)} bạn trùng tên, file có nhiều dòng hơn."

    query = _normalize_diacritic_text(name)
    scored = sorted(
        ((_token_sort_ratio(query, _normalize_diacritic_text(student.student_name)), student) for student in ordered),
        key=lambda pair: pair[0],
        reverse=True,
    )
    if scored:
        best_score, best = scored[0]
        second = scored[1][0] if len(scored) > 1 else 0
        if best_score >= FUZZY_MIN_SCORE and best_score - second >= FUZZY_MIN_MARGIN:
            return best, MATCH_FUZZY, f"Gần đúng ({best_score}%) — cần kiểm tra."
    return None, NO_MATCH, "Không tìm thấy học sinh trong lớp."


def summarize_plan(plan: Sequence[ImportPlanRow], class_size: int) -> str:
    """Một dòng tóm tắt cho hộp thoại xem trước."""

    confident = sum(1 for item in plan if item.can_apply and item.status in CONFIDENT_MATCHES)
    review = sum(1 for item in plan if item.can_apply and item.status not in CONFIDENT_MATCHES)
    unmatched = sum(1 for item in plan if not item.row_key)
    other = sum(1 for item in plan if item.row_key and not item.can_apply)
    matched_keys = {item.row_key for item in plan if item.row_key}
    missing = max(0, class_size - len(matched_keys))
    return (
        f"Khớp chắc chắn: {confident}  •  Cần kiểm tra: {review}  •  Không khớp: {unmatched}"
        f"  •  Bỏ qua (trống/trùng/lỗi): {other}  •  Học sinh trong lớp không có trong file: {missing}"
    )
