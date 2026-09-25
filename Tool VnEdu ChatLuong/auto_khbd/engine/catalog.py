"""Danh mục bài học (PPCT) và tên bài HĐTN đọc từ file Word."""

from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, TYPE_CHECKING

from ..paths import TOOL_DIR
from .analyzer.hdtn import (
    _classify_hdtn_part,
    _clean_hdtn_topic_text,
    _compact_spaces,
    _ensure_sentence_period,
    _extract_grade_from_lop_text,
    _format_hdtn_topic_heading,
)
from .parser import WeekData

if TYPE_CHECKING:  # chỉ dùng cho chú thích kiểu
    from .planner import FillOp


# ######################################################################
# Section: catalog
# ######################################################################





@dataclass
class LessonEntry:
    """1 entry: (lop, mon, phan_mon, ppct) → ten_bai."""
    lop_id: str
    mon_id: str
    phan_mon_id: str
    ppct: int
    ten_bai: str
    source: str = "filled"  # "filled" | "log" | "user"
    seen_in_weeks: list[int] = field(default_factory=list)
    # Display labels (denormalized) — tiện cho UI khỏi phải lookup từ report mỗi lần
    lop_text: str = ""
    mon_text: str = ""
    phan_mon_text: str = ""

    def key(self) -> tuple[str, str, str, int]:
        return (self.lop_id, self.mon_id, self.phan_mon_id, self.ppct)


# Priority khi merge từ nhiều nguồn
_SOURCE_PRIORITY = {"user": 3, "filled": 2, "log": 1}


class LessonCatalog:
    """Catalog tên bài.

    Sử dụng:
        cat = LessonCatalog()
        cat.ingest_weeks(weeks_data)
        cat.ingest_log(log_records)
        ten_bai = cat.lookup(lop_id, mon_id, pm_id, ppct=103)
    """

    def __init__(self):
        self._entries: dict[tuple, LessonEntry] = {}

    # -----------------------------------------------------------
    # INGEST
    # -----------------------------------------------------------

    def ingest_weeks(self, weeks: dict[int, WeekData]) -> int:
        """Ingest từ filled slots. Trả số entry mới."""
        added = 0
        for tuan in sorted(weeks.keys()):
            wd = weeks[tuan]
            for slot in wd.filled_slots:
                if not slot.lop_id or not slot.phan_mon_id or not slot.ten_bai:
                    continue
                try:
                    ppct = int(slot.ppct)
                except (ValueError, TypeError):
                    continue
                if ppct <= 0:
                    continue
                entry = LessonEntry(
                    lop_id=slot.lop_id,
                    mon_id=slot.mon_id,
                    phan_mon_id=slot.phan_mon_id,
                    ppct=ppct,
                    ten_bai=slot.ten_bai.strip(),
                    source="filled",
                    seen_in_weeks=[tuan],
                    lop_text=slot.lop_text or "",
                    mon_text=slot.mon_text or "",
                    phan_mon_text=slot.phan_mon_text or "",
                )
                if self._upsert(entry):
                    added += 1
        return added

    def ingest_log(self, log_records: Iterable[dict]) -> int:
        """Ingest từ log delete records (response của getLichBaoGiangLog).

        Mỗi record có: lop_hoc_id, mon_hoc_id, phan_mon_id, tiet_ppct, ten_bai_day, tuan
        """
        added = 0
        for rec in log_records:
            lop_id = str(rec.get("lop_hoc_id", "") or "")
            mon_id = str(rec.get("mon_hoc_id", "") or "")
            pm_id = str(rec.get("phan_mon_id", "") or "")
            ten_bai = str(rec.get("ten_bai_day", "") or "").strip()
            try:
                ppct = int(rec.get("tiet_ppct", 0) or 0)
            except (ValueError, TypeError):
                continue
            if not (lop_id and pm_id and ten_bai and ppct > 0):
                continue
            tuan = int(rec.get("tuan", 0) or 0) or None
            entry = LessonEntry(
                lop_id=lop_id, mon_id=mon_id, phan_mon_id=pm_id,
                ppct=ppct, ten_bai=ten_bai, source="log",
                seen_in_weeks=[tuan] if tuan else [],
            )
            if self._upsert(entry):
                added += 1
        return added

    def add_user_entry(self, lop_id: str, mon_id: str, phan_mon_id: str,
                       ppct: int, ten_bai: str) -> bool:
        """Override 1 entry với source 'user'."""
        entry = LessonEntry(
            lop_id=lop_id, mon_id=mon_id, phan_mon_id=phan_mon_id,
            ppct=int(ppct), ten_bai=str(ten_bai).strip(), source="user",
        )
        return self._upsert(entry)

    def _upsert(self, new: LessonEntry) -> bool:
        """Insert hoặc update theo source priority. Trả True nếu thay đổi."""
        key = new.key()
        existing = self._entries.get(key)
        if existing is None:
            self._entries[key] = new
            return True
        # Conflict: chọn source ưu tiên cao hơn
        new_pri = _SOURCE_PRIORITY.get(new.source, 0)
        ex_pri = _SOURCE_PRIORITY.get(existing.source, 0)
        if new_pri > ex_pri:
            new.seen_in_weeks = list(set(existing.seen_in_weeks + new.seen_in_weeks))
            # Carry-over labels nếu new bị thiếu
            new.lop_text = new.lop_text or existing.lop_text
            new.mon_text = new.mon_text or existing.mon_text
            new.phan_mon_text = new.phan_mon_text or existing.phan_mon_text
            self._entries[key] = new
            return True
        if new_pri == ex_pri:
            # Cùng source, merge weeks + enrich labels
            merged = list(set(existing.seen_in_weeks + new.seen_in_weeks))
            existing.seen_in_weeks = sorted([w for w in merged if w])
            if not existing.lop_text and new.lop_text:
                existing.lop_text = new.lop_text
            if not existing.mon_text and new.mon_text:
                existing.mon_text = new.mon_text
            if not existing.phan_mon_text and new.phan_mon_text:
                existing.phan_mon_text = new.phan_mon_text
            return False
        return False

    # -----------------------------------------------------------
    # LOOKUP
    # -----------------------------------------------------------

    def lookup(self, lop_id: str, mon_id: str, phan_mon_id: str, ppct: int) -> str | None:
        """Tìm tên bài cho 1 (lop, mon, pm, ppct). Trả None nếu không có."""
        key = (str(lop_id), str(mon_id), str(phan_mon_id), int(ppct))
        e = self._entries.get(key)
        return e.ten_bai if e else None

    def lookup_with_fallback(self, lop_id: str, mon_id: str, phan_mon_id: str, ppct: int) -> str | None:
        """Tìm với fallback: nếu (lop, mon, pm, ppct) không có thì thử các lớp khác cùng (mon, pm, ppct)."""
        direct = self.lookup(lop_id, mon_id, phan_mon_id, ppct)
        if direct:
            return direct
        # Fallback: cùng phan_mon + ppct, lớp khác (giả định khối tương tự)
        for k, e in self._entries.items():
            l, m, p, c = k
            if m == str(mon_id) and p == str(phan_mon_id) and c == int(ppct):
                return e.ten_bai
        return None

    def list_by_subject(self, lop_id: str, mon_id: str, phan_mon_id: str) -> list[LessonEntry]:
        """Trả list entry đã sort theo PPCT."""
        out = [e for k, e in self._entries.items()
              if k[0] == str(lop_id) and k[1] == str(mon_id) and k[2] == str(phan_mon_id)]
        return sorted(out, key=lambda e: e.ppct)

    def all_entries(self) -> list[LessonEntry]:
        return list(self._entries.values())

    # -----------------------------------------------------------
    # PERSIST
    # -----------------------------------------------------------

    def to_dict(self) -> dict:
        return {
            "version": 1,
            "entries": [
                {
                    "lop_id": e.lop_id, "mon_id": e.mon_id, "phan_mon_id": e.phan_mon_id,
                    "ppct": e.ppct, "ten_bai": e.ten_bai, "source": e.source,
                    "seen_in_weeks": e.seen_in_weeks,
                }
                for e in self._entries.values()
            ],
        }

    @classmethod
    def from_dict(cls, data: dict) -> "LessonCatalog":
        cat = cls()
        for e in (data.get("entries") or []):
            entry = LessonEntry(
                lop_id=str(e.get("lop_id", "")),
                mon_id=str(e.get("mon_id", "")),
                phan_mon_id=str(e.get("phan_mon_id", "")),
                ppct=int(e.get("ppct", 0)),
                ten_bai=str(e.get("ten_bai", "")),
                source=str(e.get("source", "user")),
                seen_in_weeks=list(e.get("seen_in_weeks", [])),
            )
            cat._upsert(entry)
        return cat

    def save_json(self, path: str | Path) -> None:
        Path(path).write_text(json.dumps(self.to_dict(), ensure_ascii=False, indent=2),
                            encoding="utf-8")

    @classmethod
    def load_json(cls, path: str | Path) -> "LessonCatalog":
        return cls.from_dict(json.loads(Path(path).read_text(encoding="utf-8")))

    def export_csv(self, path: str | Path) -> None:
        with Path(path).open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["lop_id", "mon_id", "phan_mon_id", "ppct", "ten_bai", "source", "seen_in_weeks"])
            for e in sorted(self._entries.values(), key=lambda x: (x.lop_id, x.phan_mon_id, x.ppct)):
                w.writerow([e.lop_id, e.mon_id, e.phan_mon_id, e.ppct, e.ten_bai,
                          e.source, "|".join(str(t) for t in e.seen_in_weeks)])

    def import_csv(self, path: str | Path) -> int:
        added = 0
        with Path(path).open("r", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            for row in reader:
                try:
                    e = LessonEntry(
                        lop_id=str(row.get("lop_id", "")),
                        mon_id=str(row.get("mon_id", "")),
                        phan_mon_id=str(row.get("phan_mon_id", "")),
                        ppct=int(row.get("ppct", 0)),
                        ten_bai=str(row.get("ten_bai", "")),
                        source=str(row.get("source", "user")) or "user",
                    )
                    if self._upsert(e):
                        added += 1
                except Exception:
                    continue
        return added

    def __len__(self) -> int:
        return len(self._entries)


@dataclass(frozen=True)
class HDTNLessonTitleEntry:
    """Một dòng tên bài HĐTN đọc từ kế hoạch Word."""
    ppct: int
    week: int
    topic: str
    topic_lesson: int
    detail: str = ""

    def short_title(self) -> str:
        return f"{self.topic} - Tiết {self.topic_lesson}"

    def compact_pair_title(self) -> str:
        return f"{self.topic} (tiết {self.topic_lesson})"


class HDTNLessonTitleCatalog:
    """Kho tên bài HĐTN dùng khi VnEdu không tự load tên bài."""

    DEFAULT_FILES = {
        "chu_de": Path("HĐTN") / "HDTN 6 CHỦ ĐỀ.docx",
        "shdc": Path("HĐTN") / "HDTN 6_25-26 - SHDC.docx",
        "shl": Path("HĐTN") / "HDTN 6_25-26 - SHL.docx",
    }
    GRADE8_CHU_DE_FILE = Path("HĐTN") / "25-26 PHỤ LỤC III HDTN 8.docx"

    def __init__(self, entries: dict[str, dict[int, list[HDTNLessonTitleEntry]]] | None = None):
        self._entries = entries or {
            "6:chu_de": {}, "6:shdc": {}, "6:shl": {},
            "8:chu_de": {},
        }
        self.load_errors: list[str] = []

    @classmethod
    def load_default(cls, base_dir: str | Path | None = None) -> "HDTNLessonTitleCatalog":
        """Đọc các file Word HĐTN cạnh tool, fail-soft nếu thiếu file."""
        base = Path(base_dir) if base_dir else TOOL_DIR
        cat = cls()
        try:
            from docx import Document
        except Exception as e:
            cat.load_errors.append(f"Thiếu python-docx: {e}")
            return cat

        for kind, rel_path in cls.DEFAULT_FILES.items():
            path = base / rel_path
            if not path.exists():
                cat.load_errors.append(f"Không thấy file {path}")
                continue
            try:
                cat._read_word_table(kind, Document(str(path)), grade=6)
            except Exception as e:
                cat.load_errors.append(f"Không đọc được {path.name}: {e}")

        grade8_path = base / cls.GRADE8_CHU_DE_FILE
        if grade8_path.exists():
            try:
                cat._read_grade8_chu_de_word(Document(str(grade8_path)))
            except Exception as e:
                cat.load_errors.append(f"Không đọc được {grade8_path.name}: {e}")
        return cat

    @staticmethod
    def _catalog_key(grade: int, kind: str) -> str:
        return f"{int(grade)}:{kind}"

    def _read_word_table(self, kind: str, document, grade: int) -> None:
        if len(document.tables) < 2:
            self.load_errors.append(f"lớp {grade} {kind}: file Word không có bảng phân phối")
            return
        table = document.tables[1]
        current_topic = ""
        topic_lesson = 0
        catalog_key = self._catalog_key(grade, kind)
        for row in table.rows[1:]:
            cells = [_compact_spaces(c.text) for c in row.cells]
            if not cells:
                continue
            first = cells[0]
            second = cells[1] if len(cells) > 1 else ""
            if not re.fullmatch(r"\d+", first or ""):
                topic = _clean_hdtn_topic_text(first or second)
                if topic and not topic.startswith("..."):
                    current_topic = topic
                    topic_lesson = 0
                continue
            if not current_topic:
                continue
            ppct = int(first)
            topic_lesson += 1
            week = self._extract_week(cells[3] if len(cells) > 3 else "")
            entry = HDTNLessonTitleEntry(
                ppct=ppct,
                week=week,
                topic=current_topic,
                topic_lesson=topic_lesson,
                detail=_compact_spaces(second),
            )
            self._entries.setdefault(catalog_key, {}).setdefault(ppct, []).append(entry)

    def _read_grade8_chu_de_word(self, document) -> None:
        """Đọc HĐTN 8: mỗi tuần 3 tiết, dòng giữa là HĐGD theo chủ đề."""
        if len(document.tables) < 3:
            self.load_errors.append("lớp 8 chủ đề: file Word thiếu bảng học kỳ")
            return
        week_rows: dict[int, list[tuple[int, str, str]]] = defaultdict(list)
        current_topic = ""
        for table_index in (1, 2):
            table = document.tables[table_index]
            for row in table.rows[1:]:
                cells = [_compact_spaces(c.text) for c in row.cells]
                if len(cells) < 5:
                    continue
                week = self._extract_week(cells[0])
                word_lesson = self._extract_int_cell(cells[1])
                content = cells[4]
                if not week or not word_lesson:
                    if content and "CHỦ ĐỀ" in content.upper():
                        current_topic = _format_hdtn_topic_heading(content)
                    continue
                topic_from_content = self._extract_topic_from_grade8_content(content)
                if topic_from_content:
                    current_topic = topic_from_content
                week_rows.setdefault(week, []).append((word_lesson, content, current_topic))

        catalog_key = self._catalog_key(8, "chu_de")
        topic_counts: dict[str, int] = defaultdict(int)
        for ppct, week in enumerate(sorted(week_rows), start=1):
            rows = week_rows[week]
            if len(rows) < 2:
                self.load_errors.append(f"lớp 8 chủ đề: tuần {week} thiếu dòng giữa")
                continue
            word_lesson, content, topic = rows[1]
            if not topic:
                self.load_errors.append(f"lớp 8 chủ đề: tuần {week} không xác định được chủ đề")
                continue
            topic_counts[topic] += 1
            entry = HDTNLessonTitleEntry(
                ppct=ppct,
                week=week,
                topic=topic,
                topic_lesson=topic_counts[topic],
                detail=_compact_spaces(content),
            )
            self._entries.setdefault(catalog_key, {}).setdefault(ppct, []).append(entry)

    @staticmethod
    def _extract_week(text: str) -> int:
        m = re.search(r"\d+", str(text or ""))
        return int(m.group(0)) if m else 0

    @staticmethod
    def _extract_int_cell(text: str) -> int:
        s = _compact_spaces(text)
        return int(s) if re.fullmatch(r"\d+", s or "") else 0

    @staticmethod
    def _extract_topic_from_grade8_content(content: str) -> str:
        c = _compact_spaces(content)
        m = re.search(
            r"(Chủ\s*đề\s*\d+\s*[:\.]?\s*.*?)(?=\s*-\s*NV\d|\s+-NV\d|\s+NV\d|$)",
            c,
            flags=re.I,
        )
        return _format_hdtn_topic_heading(m.group(1).rstrip(" -.;")) if m else ""

    def lookup_for_op(self, op: "FillOp") -> str | None:
        """Tìm tên bài theo khối/phân môn/PPCT."""
        if not op or not op.ppct:
            return None
        grade = _extract_grade_from_lop_text(op.lop_text)
        if grade not in (6, 8):
            return None
        kind = _classify_hdtn_part(op.mon_text, op.phan_mon_text)
        if grade == 8 and kind != "chu_de":
            return None
        catalog_key = self._catalog_key(grade, kind)
        if catalog_key not in self._entries:
            return None
        matches = self._entries.get(catalog_key, {}).get(int(op.ppct), [])
        if not matches:
            return None
        if len(matches) > 1:
            return " - ".join(e.compact_pair_title() for e in matches)
        entry = matches[0]
        if kind == "shdc":
            title = f"{entry.short_title()}: {_compact_spaces(entry.detail)}"
            return _ensure_sentence_period(title)
        return entry.short_title()


_HDTN_TITLE_CATALOG_CACHE: HDTNLessonTitleCatalog | None = None


def get_hdtn_lesson_title_catalog() -> HDTNLessonTitleCatalog:
    """Cache kho tên bài HĐTN để mỗi phiên nhập không đọc Word lặp lại."""
    global _HDTN_TITLE_CATALOG_CACHE
    if _HDTN_TITLE_CATALOG_CACHE is None:
        _HDTN_TITLE_CATALOG_CACHE = HDTNLessonTitleCatalog.load_default()
    return _HDTN_TITLE_CATALOG_CACHE
