"""KHBD Logic Engine — kiểm tra / sửa / so sánh PPCT trên bản sao lưu."""

from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable

from ..analyzer.hdtn import normalize_text_loose
from ..analyzer.models import (
    _SEVERITY_ORDER,
    SEVERITY_ERROR,
    SEVERITY_INFO,
    SEVERITY_WARNING,
    TKB_EVENT_NGHI,
    TT_BINH_THUONG,
    TT_CHEN_LICH,
    TT_COUNTED,
    TT_DAY_BU,
    TT_DAY_CHUNG,
    TT_DAY_THAY,
    TT_NGHI,
)
from ..profile.models import TKBScheduleEvent
from .models import (
    _TT_VALUE_TO_TEXT_CANONICAL,
    BACKUP_FILE_EXT,
    BackupFile,
    BackupMetadata,
    BackupSchemaError,
    BackupSlot,
    BackupWeek,
    load_backup_json,
)


# ######################################################################
# KHBD Logic Engine — audit / repair / diff trên backup hoặc Excel review
# ######################################################################

KHBD_ISSUE_ACTIVE_DUP = "active_duplicate"


KHBD_ISSUE_PPCT_GAP = "ppct_gap"


KHBD_ISSUE_PPCT_ORDER = "ppct_order"


KHBD_ISSUE_OVER_CAP = "over_cap"


KHBD_ISSUE_NGHI_NO_MAKEUP = "nghi_no_makeup"


KHBD_ISSUE_TITLE_CONFLICT = "title_conflict"


@dataclass
class KHBDAuditIssue:
    """Một lỗi/cảnh báo dữ liệu KHBD trước khi import hoặc ghi web."""
    severity: str
    kind: str
    tuan: int
    row_key: str
    group_key: str
    label: str
    ppct: str = ""
    message: str = ""


@dataclass
class KHBDRepairAction:
    """Một đề xuất sửa offline, chỉ áp dụng khi user duyệt preview."""
    tuan: int
    row_key: str
    group_key: str
    label: str
    current_ppct: str = ""
    proposed_ppct: str = ""
    current_ten_bai: str = ""
    proposed_ten_bai: str = ""
    current_status: str = ""
    proposed_status: str = ""
    note: str = ""
    skip: bool = False

    @property
    def needs_change(self) -> bool:
        return (
            str(self.current_ppct or "") != str(self.proposed_ppct or "")
            or str(self.current_ten_bai or "") != str(self.proposed_ten_bai or "")
            or (
                bool(self.proposed_status)
                and str(self.current_status or "") != str(self.proposed_status or "")
            )
        )


@dataclass
class KHBDAuditResult:
    """Kết quả audit tổng hợp của KHBD Logic Engine."""
    issues: list[KHBDAuditIssue] = field(default_factory=list)
    groups_checked: int = 0
    slots_checked: int = 0

    @property
    def error_count(self) -> int:
        return sum(1 for i in self.issues if i.severity == SEVERITY_ERROR)

    @property
    def warning_count(self) -> int:
        return sum(1 for i in self.issues if i.severity == SEVERITY_WARNING)

    @property
    def blocker_count(self) -> int:
        return self.error_count


@dataclass
class KHBDDiffRow:
    """Một dòng khác biệt giữa 2 nguồn backup/review."""
    tuan: int
    row_key: str
    field: str
    left_value: str
    right_value: str
    label: str
    severity: str = SEVERITY_INFO


@dataclass
class KHBDHolidayRequest:
    """Yêu cầu dựng ngày Nghỉ + Dạy bù trên backup offline."""
    holiday_tuan: int
    thu: int
    buoi_idx: int = 0        # 0 = cả ngày, 1 = sáng, 2 = chiều
    makeup_tuan: int = 36
    note: str = ""
    include_unmade_nghi_in_week: bool = True
    auto_resequence: bool = True


class KHBDLogicEngine:
    """Audit và sửa dữ liệu KHBD offline theo luật PPCT đã thống nhất.

    Engine này không thao tác web. Nó chỉ làm việc trên `BackupFile` hoặc
    dữ liệu Excel review đã được quy đổi về `BackupFile`.
    """

    TITLE_FIELDS = ("ppct", "ten_bai", "ghi_chu", "id_giao_an")

    @staticmethod
    def clone_backup(bf: BackupFile) -> BackupFile:
        return BackupFile.from_dict(deepcopy(bf.to_dict()))

    @staticmethod
    def group_key(slot: BackupSlot) -> str:
        return f"{slot.lop_id}|{slot.mon_id}|{slot.phan_mon_id}"

    @staticmethod
    def group_label(slot: BackupSlot) -> str:
        return f"{slot.lop_text} / {slot.mon_text} / {slot.phan_mon_text}"

    @staticmethod
    def schedule_key(tuan: int, slot: BackupSlot) -> tuple[int, int, int, int, str]:
        return (
            int(tuan or 0), int(slot.thu or 0), int(slot.buoi_idx or 0),
            int(slot.tiet_idx or 0), str(slot.row_key or ""),
        )

    @staticmethod
    def ppct_int(slot: BackupSlot) -> int:
        try:
            return int(str(slot.ppct or "").strip())
        except (TypeError, ValueError):
            return 0

    @staticmethod
    def _status_code(slot: BackupSlot) -> str:
        code = str(slot.trang_thai or "").strip()
        if code:
            return code
        text = normalize_text_loose(slot.trang_thai_text or "")
        if text == normalize_text_loose("Nghỉ"):
            return TT_NGHI
        if text == normalize_text_loose("Dạy bù"):
            return TT_DAY_BU
        if text == normalize_text_loose("Dạy thay"):
            return TT_DAY_THAY
        if text == normalize_text_loose("Chèn lịch"):
            return TT_CHEN_LICH
        if text == normalize_text_loose("Dạy chung"):
            return TT_DAY_CHUNG
        if text == normalize_text_loose("Bình thường"):
            return TT_BINH_THUONG
        return code

    @classmethod
    def is_nghi(cls, slot: BackupSlot) -> bool:
        return cls._status_code(slot) == TT_NGHI

    @classmethod
    def is_active_counted(cls, slot: BackupSlot) -> bool:
        if cls.ppct_int(slot) <= 0:
            return False
        code = cls._status_code(slot)
        if code == TT_NGHI:
            return False
        if not code:
            return True
        return code in TT_COUNTED

    @staticmethod
    def subject_cap(slot: BackupSlot) -> int | None:
        mon = normalize_text_loose(slot.mon_text or "")
        pm = normalize_text_loose(slot.phan_mon_text or "")
        if mon == normalize_text_loose("Ngoại ngữ"):
            return 105
        if pm in {
            normalize_text_loose("Sinh hoạt lớp"),
            normalize_text_loose("Sinh hoạt dưới cờ"),
        }:
            return 35
        return None

    @classmethod
    def iter_slots(cls, bf: BackupFile) -> Iterable[tuple[int, BackupSlot]]:
        for week in sorted(bf.weeks, key=lambda w: int(w.tuan or 0)):
            for slot in sorted(
                week.slots,
                key=lambda s: cls.schedule_key(int(week.tuan), s),
            ):
                yield int(week.tuan), slot

    @classmethod
    def group_slots(
        cls, bf: BackupFile,
    ) -> dict[str, list[tuple[int, BackupSlot]]]:
        out: dict[str, list[tuple[int, BackupSlot]]] = defaultdict(list)
        for tuan, slot in cls.iter_slots(bf):
            if not slot.has_lop or cls.ppct_int(slot) <= 0:
                continue
            out[cls.group_key(slot)].append((tuan, slot))
        return out

    @classmethod
    def _catalog_by_ppct(
        cls, items: list[tuple[int, BackupSlot]],
    ) -> dict[int, str]:
        catalog: dict[int, str] = {}
        for _tuan, slot in sorted(
            items, key=lambda x: cls.schedule_key(x[0], x[1]),
        ):
            ppct = cls.ppct_int(slot)
            if ppct <= 0:
                continue
            ten_bai = str(slot.ten_bai or "").strip()
            if ten_bai and ppct not in catalog:
                catalog[ppct] = ten_bai
        return catalog

    @classmethod
    def analyze_backup(cls, bf: BackupFile) -> KHBDAuditResult:
        result = KHBDAuditResult(slots_checked=bf.total_slots)
        groups = cls.group_slots(bf)
        result.groups_checked = len(groups)
        for group_key, items in sorted(groups.items()):
            label = cls.group_label(items[0][1])
            active = [
                (t, s) for t, s in items if cls.is_active_counted(s)
            ]
            nghi = [
                (t, s) for t, s in items
                if cls.is_nghi(s) and cls.ppct_int(s) > 0
            ]
            active_sorted = sorted(
                active, key=lambda x: cls.schedule_key(x[0], x[1]),
            )
            active_ppcts = [cls.ppct_int(s) for _t, s in active_sorted]
            active_by_ppct: dict[int, list[tuple[int, BackupSlot]]] = defaultdict(list)
            for tuan, slot in active_sorted:
                active_by_ppct[cls.ppct_int(slot)].append((tuan, slot))

            for ppct, pp_items in sorted(active_by_ppct.items()):
                if len(pp_items) <= 1:
                    continue
                first_tuan, first_slot = pp_items[0]
                rows = ", ".join(
                    f"T{t} {s.row_key}" for t, s in pp_items[:6]
                )
                result.issues.append(KHBDAuditIssue(
                    severity=SEVERITY_ERROR,
                    kind=KHBD_ISSUE_ACTIVE_DUP,
                    tuan=first_tuan,
                    row_key=first_slot.row_key,
                    group_key=group_key,
                    label=label,
                    ppct=str(ppct),
                    message=f"PPCT active {ppct} bị trùng: {rows}",
                ))

            if active_ppcts:
                existing = set(active_ppcts)
                for miss in range(min(active_ppcts), max(active_ppcts) + 1):
                    if miss not in existing:
                        first_tuan, first_slot = active_sorted[0]
                        result.issues.append(KHBDAuditIssue(
                            severity=SEVERITY_ERROR,
                            kind=KHBD_ISSUE_PPCT_GAP,
                            tuan=first_tuan,
                            row_key=first_slot.row_key,
                            group_key=group_key,
                            label=label,
                            ppct=str(miss),
                            message=f"Thiếu active PPCT {miss}.",
                        ))

                for (prev_t, prev_s), (cur_t, cur_s) in zip(
                    active_sorted, active_sorted[1:],
                ):
                    prev_ppct = cls.ppct_int(prev_s)
                    cur_ppct = cls.ppct_int(cur_s)
                    if cur_ppct - prev_ppct != 1:
                        result.issues.append(KHBDAuditIssue(
                            severity=SEVERITY_ERROR,
                            kind=KHBD_ISSUE_PPCT_ORDER,
                            tuan=cur_t,
                            row_key=cur_s.row_key,
                            group_key=group_key,
                            label=label,
                            ppct=str(cur_ppct),
                            message=(
                                f"Thứ tự PPCT sai: T{prev_t} {prev_s.row_key} "
                                f"PPCT {prev_ppct} -> T{cur_t} {cur_s.row_key} "
                                f"PPCT {cur_ppct}."
                            ),
                        ))

                cap = cls.subject_cap(active_sorted[0][1])
                if cap is not None:
                    for tuan, slot in active_sorted:
                        ppct = cls.ppct_int(slot)
                        if ppct > cap:
                            result.issues.append(KHBDAuditIssue(
                                severity=SEVERITY_ERROR,
                                kind=KHBD_ISSUE_OVER_CAP,
                                tuan=tuan,
                                row_key=slot.row_key,
                                group_key=group_key,
                                label=label,
                                ppct=str(ppct),
                                message=f"PPCT {ppct} vượt trần {cap}.",
                            ))

            for tuan, slot in nghi:
                ppct = cls.ppct_int(slot)
                if ppct not in active_by_ppct:
                    result.issues.append(KHBDAuditIssue(
                        severity=SEVERITY_ERROR,
                        kind=KHBD_ISSUE_NGHI_NO_MAKEUP,
                        tuan=tuan,
                        row_key=slot.row_key,
                        group_key=group_key,
                        label=label,
                        ppct=str(ppct),
                        message=(
                            f"Tiết Nghỉ PPCT {ppct} chưa có tiết active/dạy bù "
                            "tương ứng."
                        ),
                    ))

            titles_by_ppct: dict[int, set[str]] = defaultdict(set)
            for _tuan, slot in items:
                ppct = cls.ppct_int(slot)
                title = str(slot.ten_bai or "").strip()
                if ppct > 0 and title:
                    titles_by_ppct[ppct].add(title)
            for ppct, titles in sorted(titles_by_ppct.items()):
                if len(titles) <= 1:
                    continue
                first_tuan, first_slot = items[0]
                sample = "; ".join(sorted(titles)[:3])
                result.issues.append(KHBDAuditIssue(
                    severity=SEVERITY_WARNING,
                    kind=KHBD_ISSUE_TITLE_CONFLICT,
                    tuan=first_tuan,
                    row_key=first_slot.row_key,
                    group_key=group_key,
                    label=label,
                    ppct=str(ppct),
                    message=(
                        f"Cùng PPCT {ppct} có nhiều tên bài khác nhau: {sample}"
                    ),
                ))
        result.issues.sort(key=lambda i: (
            _SEVERITY_ORDER.get(i.severity, 9), i.kind, i.label, i.tuan,
            i.row_key, i.ppct,
        ))
        return result

    @classmethod
    def build_sequential_repairs(
        cls,
        bf: BackupFile,
        group_keys: set[str] | None = None,
    ) -> list[KHBDRepairAction]:
        actions: list[KHBDRepairAction] = []
        groups = cls.group_slots(bf)
        for group_key, items in sorted(groups.items()):
            if group_keys is not None and group_key not in group_keys:
                continue
            label = cls.group_label(items[0][1])
            active = sorted(
                [(t, s) for t, s in items if cls.is_active_counted(s)],
                key=lambda x: cls.schedule_key(x[0], x[1]),
            )
            if not active:
                continue
            current_ppcts = [cls.ppct_int(s) for _t, s in active]
            start_ppct = min(current_ppcts)
            catalog = cls._catalog_by_ppct(items)
            for idx, (tuan, slot) in enumerate(active):
                proposed_ppct = start_ppct + idx
                proposed_ten_bai = catalog.get(
                    proposed_ppct, str(slot.ten_bai or ""),
                )
                action = KHBDRepairAction(
                    tuan=tuan,
                    row_key=slot.row_key,
                    group_key=group_key,
                    label=label,
                    current_ppct=str(slot.ppct or ""),
                    proposed_ppct=str(proposed_ppct),
                    current_ten_bai=str(slot.ten_bai or ""),
                    proposed_ten_bai=proposed_ten_bai,
                    current_status=slot.trang_thai_text or slot.trang_thai,
                    note="Kéo PPCT/Tên bài theo chuỗi active liên tục.",
                )
                if action.needs_change:
                    actions.append(action)
        return actions

    @classmethod
    def apply_actions(
        cls, bf: BackupFile, actions: list[KHBDRepairAction],
    ) -> int:
        by_key: dict[tuple[int, str], BackupSlot] = {}
        for tuan, slot in cls.iter_slots(bf):
            by_key[(int(tuan), str(slot.row_key))] = slot
        changed = 0
        for action in actions:
            if action.skip:
                continue
            slot = by_key.get((int(action.tuan), str(action.row_key)))
            if slot is None:
                continue
            before = (
                slot.ppct, slot.ten_bai, slot.trang_thai,
                slot.trang_thai_text,
            )
            if action.proposed_ppct:
                slot.ppct = str(action.proposed_ppct)
            slot.ten_bai = str(action.proposed_ten_bai or "")
            if action.proposed_status:
                slot.trang_thai_text = str(action.proposed_status)
                slot.trang_thai = (
                    TT_NGHI if action.proposed_status == "Nghỉ"
                    else TT_DAY_BU if action.proposed_status == "Dạy bù"
                    else slot.trang_thai
                )
            after = (
                slot.ppct, slot.ten_bai, slot.trang_thai,
                slot.trang_thai_text,
            )
            if before != after:
                changed += 1
        return changed

    @classmethod
    def apply_holiday(
        cls, bf: BackupFile, request: KHBDHolidayRequest,
    ) -> list[KHBDRepairAction]:
        actions: list[KHBDRepairAction] = []
        week = bf.week_by_num(request.holiday_tuan)
        makeup_week = bf.week_by_num(request.makeup_tuan)
        if week is None:
            raise ValueError(f"Không có tuần {request.holiday_tuan} trong file.")
        if makeup_week is None:
            makeup_week = BackupWeek(
                tuan=int(request.makeup_tuan),
                fetched_at=datetime.now().isoformat(timespec="seconds"),
                slots=[],
            )
            bf.weeks.append(makeup_week)
            bf.weeks.sort(key=lambda w: int(w.tuan or 0))

        affected_groups: set[str] = set()
        cancelled: list[BackupSlot] = []
        for slot in week.slots:
            if int(slot.thu or 0) != int(request.thu):
                continue
            if request.buoi_idx and int(slot.buoi_idx or 0) != int(request.buoi_idx):
                continue
            if not slot.has_lop or cls.ppct_int(slot) <= 0:
                continue
            if not cls.is_nghi(slot):
                actions.append(KHBDRepairAction(
                    tuan=week.tuan,
                    row_key=slot.row_key,
                    group_key=cls.group_key(slot),
                    label=cls.group_label(slot),
                    current_ppct=str(slot.ppct or ""),
                    proposed_ppct=str(slot.ppct or ""),
                    current_ten_bai=str(slot.ten_bai or ""),
                    proposed_ten_bai=str(slot.ten_bai or ""),
                    current_status=slot.trang_thai_text or slot.trang_thai,
                    proposed_status="Nghỉ",
                    note=request.note or "Ngày nghỉ được khai báo trong wizard.",
                ))
                slot.trang_thai = TT_NGHI
                slot.trang_thai_text = "Nghỉ"
                cancelled.append(slot)
                affected_groups.add(cls.group_key(slot))

        if request.include_unmade_nghi_in_week:
            analysis = cls.analyze_backup(bf)
            missing_keys = {
                (i.tuan, i.row_key)
                for i in analysis.issues
                if i.kind == KHBD_ISSUE_NGHI_NO_MAKEUP
                and int(i.tuan) == int(request.holiday_tuan)
            }
            for slot in week.slots:
                if (int(week.tuan), slot.row_key) not in missing_keys:
                    continue
                if slot not in cancelled:
                    cancelled.append(slot)
                    affected_groups.add(cls.group_key(slot))

        existing_by_row = {str(s.row_key): s for s in makeup_week.slots}
        for source in cancelled:
            target = existing_by_row.get(str(source.row_key))
            if target is not None and target.has_lop and cls.group_key(target) != cls.group_key(source):
                actions.append(KHBDRepairAction(
                    tuan=makeup_week.tuan,
                    row_key=source.row_key,
                    group_key=cls.group_key(source),
                    label=cls.group_label(source),
                    current_ppct="",
                    proposed_ppct=str(source.ppct or ""),
                    current_ten_bai="",
                    proposed_ten_bai=str(source.ten_bai or ""),
                    current_status="",
                    proposed_status="Dạy bù",
                    note=(
                        f"Bỏ qua: tuần {makeup_week.tuan} ô {source.row_key} "
                        "đã có tiết khác."
                    ),
                    skip=True,
                ))
                continue
            if target is None:
                target = deepcopy(source)
                makeup_week.slots.append(target)
                existing_by_row[str(target.row_key)] = target
            else:
                for attr in (
                    "lop_id", "lop_text", "khoi", "mon_id", "mon_text",
                    "mon_dmonid", "phan_mon_id", "phan_mon_text",
                    "ppct", "ten_bai", "ghi_chu", "id_giao_an",
                ):
                    setattr(target, attr, getattr(source, attr))
            target.trang_thai = TT_DAY_BU
            target.trang_thai_text = "Dạy bù"
            affected_groups.add(cls.group_key(target))
            actions.append(KHBDRepairAction(
                tuan=makeup_week.tuan,
                row_key=target.row_key,
                group_key=cls.group_key(target),
                label=cls.group_label(target),
                current_ppct="",
                proposed_ppct=str(target.ppct or ""),
                current_ten_bai="",
                proposed_ten_bai=str(target.ten_bai or ""),
                current_status="",
                proposed_status="Dạy bù",
                note="Tạo/cập nhật tiết Dạy bù ở tuần bù.",
            ))

        makeup_week.slots.sort(
            key=lambda s: cls.schedule_key(int(makeup_week.tuan), s),
        )
        if request.auto_resequence and affected_groups:
            seq_actions = cls.build_sequential_repairs(
                bf, group_keys=affected_groups,
            )
            cls.apply_actions(bf, seq_actions)
            actions.extend(seq_actions)
        return [a for a in actions if a.needs_change or a.note]

    @classmethod
    def apply_schedule_events_to_backup(
        cls,
        bf: BackupFile,
        events: list[TKBScheduleEvent],
    ) -> list[KHBDRepairAction]:
        """Áp dụng event Nghỉ/Dạy bù mới vào backup offline rồi kéo lại PPCT.

        Dùng cùng logic counted với audit: Nghỉ không tính PPCT, Dạy bù có
        tính PPCT. Hàm này không ghi file; caller tự preview và lưu JSON.
        """
        actions: list[KHBDRepairAction] = []
        affected_groups: set[str] = set()

        def _week(tuan: int) -> BackupWeek:
            week = bf.week_by_num(tuan)
            if week is None:
                week = BackupWeek(
                    tuan=int(tuan),
                    fetched_at=datetime.now().isoformat(timespec="seconds"),
                    slots=[],
                )
                bf.weeks.append(week)
                bf.weeks.sort(key=lambda w: int(w.tuan or 0))
            return week

        def _slot_by_row(week: BackupWeek, row_key: str) -> BackupSlot | None:
            for slot in week.slots:
                if str(slot.row_key) == str(row_key):
                    return slot
            return None

        def _event_group_label(event: TKBScheduleEvent) -> str:
            return (
                f"{event.source_lop_text} / {event.source_mon_text} / "
                f"{event.source_phan_mon_text}"
            )

        def _max_active_ppct(group_key: str) -> int:
            values = [
                cls.ppct_int(slot)
                for _tuan, slot in cls.iter_slots(bf)
                if cls.group_key(slot) == group_key and cls.is_active_counted(slot)
            ]
            return max(values) if values else 0

        for event in sorted(
            [e for e in events if getattr(e, "enabled", True)],
            key=lambda e: e.position_tuple(),
        ):
            week = _week(event.tuan)
            target = _slot_by_row(week, event.row_key)
            group_key = event.group_key
            label = _event_group_label(event)

            if event.kind == TKB_EVENT_NGHI:
                if target is None or not target.has_lop:
                    actions.append(KHBDRepairAction(
                        tuan=event.tuan,
                        row_key=event.row_key,
                        group_key=group_key,
                        label=label,
                        note="Bỏ qua: không tìm thấy tiết TKB để đánh dấu Nghỉ.",
                        skip=True,
                    ))
                    continue
                if cls.group_key(target) != group_key:
                    actions.append(KHBDRepairAction(
                        tuan=event.tuan,
                        row_key=event.row_key,
                        group_key=group_key,
                        label=label,
                        current_status=target.trang_thai_text or target.trang_thai,
                        note="Bỏ qua: ô Nghỉ không còn đúng lớp/môn/phân môn nguồn.",
                        skip=True,
                    ))
                    continue
                before_status = target.trang_thai_text or target.trang_thai
                target.trang_thai = TT_NGHI
                target.trang_thai_text = "Nghỉ"
                if event.note:
                    target.ghi_chu = event.note
                affected_groups.add(group_key)
                actions.append(KHBDRepairAction(
                    tuan=event.tuan,
                    row_key=event.row_key,
                    group_key=group_key,
                    label=cls.group_label(target),
                    current_ppct=str(target.ppct or ""),
                    proposed_ppct=str(target.ppct or ""),
                    current_ten_bai=str(target.ten_bai or ""),
                    proposed_ten_bai=str(target.ten_bai or ""),
                    current_status=before_status,
                    proposed_status="Nghỉ",
                    note=event.note or "Đánh dấu Nghỉ từ hồ sơ TKB.",
                ))
                continue

            old_group = cls.group_key(target) if (target is not None and target.has_lop) else ""
            if target is not None and target.has_lop and old_group != group_key:
                if not event.replace_normal_slot:
                    actions.append(KHBDRepairAction(
                        tuan=event.tuan,
                        row_key=event.row_key,
                        group_key=group_key,
                        label=label,
                        current_status=target.trang_thai_text or target.trang_thai,
                        note="Bỏ qua: vị trí Dạy bù đang có tiết khác và chưa cho ghi đè.",
                        skip=True,
                    ))
                    continue
                affected_groups.add(old_group)

            if target is None:
                target = BackupSlot(
                    row_key=event.row_key,
                    thu=event.thu,
                    buoi_idx=event.buoi,
                    tiet_idx=event.tiet,
                    tiet_tkb=str(event.tiet),
                )
                week.slots.append(target)

            before = (
                target.ppct, target.ten_bai,
                target.trang_thai, target.trang_thai_text,
            )
            target.lop_id = event.source_lop_id
            target.lop_text = event.source_lop_text
            target.mon_id = event.source_mon_id
            target.mon_text = event.source_mon_text
            target.phan_mon_id = event.source_phan_mon_id
            target.phan_mon_text = event.source_phan_mon_text
            target.ppct = str(_max_active_ppct(group_key) + 1)
            target.ten_bai = ""
            target.ghi_chu = event.note or "Dạy bù"
            target.trang_thai = TT_DAY_BU
            target.trang_thai_text = "Dạy bù"
            affected_groups.add(group_key)
            actions.append(KHBDRepairAction(
                tuan=event.tuan,
                row_key=event.row_key,
                group_key=group_key,
                label=label,
                current_ppct=str(before[0] or ""),
                proposed_ppct=str(target.ppct or ""),
                current_ten_bai=str(before[1] or ""),
                proposed_ten_bai=str(target.ten_bai or ""),
                current_status=str(before[3] or before[2] or ""),
                proposed_status="Dạy bù",
                note=event.note or "Tạo/cập nhật Dạy bù từ hồ sơ TKB.",
            ))

        for week in bf.weeks:
            week.slots.sort(key=lambda s: cls.schedule_key(int(week.tuan), s))
        if affected_groups:
            seq_actions = cls.build_sequential_repairs(
                bf, group_keys=affected_groups,
            )
            cls.apply_actions(bf, seq_actions)
            actions.extend(seq_actions)
        return [a for a in actions if a.needs_change or a.note]

    @classmethod
    def diff_backups(
        cls, left: BackupFile, right: BackupFile,
    ) -> list[KHBDDiffRow]:
        def build_map(bf: BackupFile) -> dict[tuple[int, str], BackupSlot]:
            return {
                (t, s.row_key): s
                for t, s in cls.iter_slots(bf)
            }

        left_map = build_map(left)
        right_map = build_map(right)
        rows: list[KHBDDiffRow] = []
        for key in sorted(set(left_map) | set(right_map)):
            tuan, row_key = key
            lslot = left_map.get(key)
            rslot = right_map.get(key)
            label = cls.group_label(rslot or lslot) if (rslot or lslot) else ""
            if lslot is None:
                rows.append(KHBDDiffRow(
                    tuan=tuan, row_key=row_key, field="slot",
                    left_value="(không có)", right_value="có thêm",
                    label=label, severity=SEVERITY_WARNING,
                ))
                continue
            if rslot is None:
                rows.append(KHBDDiffRow(
                    tuan=tuan, row_key=row_key, field="slot",
                    left_value="có", right_value="(đã bỏ)",
                    label=label, severity=SEVERITY_WARNING,
                ))
                continue
            for field_name in (
                "lop_text", "mon_text", "phan_mon_text", "ppct",
                "ten_bai", "trang_thai_text", "ghi_chu",
            ):
                left_value = str(getattr(lslot, field_name, "") or "")
                right_value = str(getattr(rslot, field_name, "") or "")
                if left_value == right_value:
                    continue
                severity = (
                    SEVERITY_ERROR
                    if field_name in {"ppct", "trang_thai_text"}
                    else SEVERITY_WARNING
                )
                rows.append(KHBDDiffRow(
                    tuan=tuan,
                    row_key=row_key,
                    field=field_name,
                    left_value=left_value,
                    right_value=right_value,
                    label=label,
                    severity=severity,
                ))
        return rows


def load_khbd_review_file(path: str | Path) -> BackupFile:
    """Load backup JSON hoặc workbook review có sheet `Du_lieu_goc`."""
    p = Path(path)
    if p.suffix.lower() == ".json" or p.name.endswith(BACKUP_FILE_EXT):
        return load_backup_json(p)
    if p.suffix.lower() != ".xlsx":
        raise BackupSchemaError("Chỉ hỗ trợ .khdh-backup.json hoặc .xlsx.")
    try:
        from openpyxl import load_workbook
    except Exception as e:
        raise BackupSchemaError(f"Thiếu openpyxl để đọc Excel: {e}")
    try:
        wb = load_workbook(p, read_only=True, data_only=False)
    except Exception as e:
        raise BackupSchemaError(f"Không mở được Excel: {e}")
    if "Du_lieu_goc" not in wb.sheetnames:
        raise BackupSchemaError("Excel không có sheet Du_lieu_goc.")
    ws = wb["Du_lieu_goc"]
    headers: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        value = ws.cell(1, c).value
        if value:
            headers[str(value).strip()] = c

    def cell(row: int, name: str, default: str = "") -> str:
        c = headers.get(name)
        if not c:
            return default
        value = ws.cell(row, c).value
        return "" if value is None else str(value)

    weeks_map: dict[int, BackupWeek] = {}
    for r in range(2, ws.max_row + 1):
        tuan_raw = cell(r, "tuan")
        row_key = cell(r, "row_key")
        if not tuan_raw or not row_key:
            continue
        try:
            tuan = int(float(tuan_raw))
        except ValueError:
            continue
        status_code = cell(r, "trang_thai")
        status_text = cell(r, "trang_thai_text")
        if not status_text:
            status_text = _TT_VALUE_TO_TEXT_CANONICAL.get(status_code, "")
        slot = BackupSlot(
            row_key=row_key,
            thu=int(float(cell(r, "thu", "0") or 0)),
            buoi_idx=int(float(cell(r, "buoi_idx", "0") or 0)),
            tiet_idx=int(float(cell(r, "tiet_idx", "0") or 0)),
            tiet_tkb=cell(r, "tiet_tkb"),
            lop_id=cell(r, "lop_id"),
            lop_text=cell(r, "lop_text"),
            khoi=cell(r, "khoi"),
            mon_id=cell(r, "mon_id"),
            mon_text=cell(r, "mon_text"),
            mon_dmonid=cell(r, "mon_dmonid"),
            phan_mon_id=cell(r, "phan_mon_id"),
            phan_mon_text=cell(r, "phan_mon_text"),
            ppct=cell(r, "ppct"),
            ten_bai=cell(r, "ten_bai"),
            ghi_chu=cell(r, "ghi_chu"),
            trang_thai=status_code,
            trang_thai_text=status_text,
            id_giao_an=cell(r, "id_giao_an"),
        )
        if not slot.has_lop:
            continue
        week = weeks_map.setdefault(tuan, BackupWeek(tuan=tuan, fetched_at=""))
        week.slots.append(slot)
    weeks = [
        BackupWeek(
            tuan=w.tuan,
            fetched_at=w.fetched_at,
            slots=sorted(
                w.slots,
                key=lambda s: KHBDLogicEngine.schedule_key(w.tuan, s),
            ),
        )
        for w in sorted(weeks_map.values(), key=lambda x: int(x.tuan))
    ]
    return BackupFile(
        metadata=BackupMetadata(
            created_at=datetime.now().isoformat(timespec="seconds"),
            tool_version="auto_khbd_pro_excel_review",
            tuan_from=min(weeks_map) if weeks_map else 1,
            tuan_to=max(weeks_map) if weeks_map else 36,
            note=f"Loaded from Excel review: {p.name}",
        ),
        weeks=weeks,
    )
